from quantities import ms, Hz, uV, s
import numpy as np
import matplotlib.pyplot as plt
import neo
from neo.io.nixio import NixIO
from neo.core import AnalogSignal, SpikeTrain, IrregularlySampledSignal
from ceed.analysis import CeedDataReader
from ceed.function import CeedFuncRef
import pandas as pd
from itertools import compress
from openpyxl import load_workbook

"""(I think) this gets the time course of a given stimulus """
def get_stimulus_signal(reader, exp, shape="enclosed", led="blue", returnas="percent"):

    reader.load_application_data()
    reader.load_experiment(exp)
    shape_intensities = np.array(reader.shapes_intensity[shape])
    alignment = np.array(reader.electrode_intensity_alignment)

    if led == "red":
        index = 0
    if led == "green":
        index = 1
    if led == "blue":
        index = 2
    intensity = shape_intensities[:-1, index]

    if returnas == "percent":
        intensity = intensity * 100
    if returnas == "norm":
        max = np.max(intensity)
        intensity = intensity / max

    fs = reader.electrodes_metadata['A4']['sampling_frequency']*Hz  # arbitrary electrode
    period = 1./fs
    times = [x * period for x in alignment]
    times = np.array(times)
    stimulus = IrregularlySampledSignal(times, intensity, units='percent', time_units='s')
    return stimulus


def get_all_exps(file):
    reader = CeedDataReader(file)
    reader.open_h5()
    reader.load_application_data()

    exps = reader.experiments_in_file

    reader.load_mcs_data()
    fs = reader.electrodes_metadata['A4']['sampling_frequency'] * Hz
    period = (1. / fs).rescale(s)
    led = ['red', 'green', 'blue']

    columns = ['experiment', 'stage', 'substage', 'function', 'intensity', 'pattern', 'color', 't_start', 't_stop',
               'signal']
    stim_df = pd.DataFrame(index=range(100), columns=columns)

    i = 0
    for j, exp in enumerate(exps):

        reader.load_experiment(exp)
        stage_dict = reader.stage_factory.stage_names
        stage = stage_dict[reader.experiment_stage_name]
        substages = stage.stages

        if len(substages) == 0:
            stim_df.loc[i, 'experiment'] = str(exp)
            stim_df.loc[i, 'stage'] = reader.experiment_stage_name
            stim_df.loc[i, 'substage'] = "None"
            if len(stage.shapes) == 0:
                print("Skipping experiment #" + str(exp) + "; no shapes were found within the stage.")
                continue
            patterns = [x.name for x in stage.shapes]
            stim_df.loc[i, 'pattern'] = tuple(patterns)

            if type(stage.functions[0]) == CeedFuncRef:
                functions = [x.func.name for x in stage.functions]
            else:
                functions = [x.name for x in stage.functions]
            stim_df.loc[i, 'function'] = functions

            color_mask = [stage.color_r, stage.color_g, stage.color_b]
            colors = list(compress(led, color_mask))
            stim_df.loc[i, 'color'] = colors

            alignment = np.array(reader.electrode_intensity_alignment)
            if len(alignment.shape) == 0:
                stim_df.loc[i, 't_start'] = np.nan
                stim_df.loc[i, 't_stop'] = np.nan
                stim_df.loc[i, 'signal'] = [np.nan]
                stim_df.loc[i, 'intensity'] = np.nan
            else:
                times = [round(x * period, 5) for x in alignment]
                stim_df.loc[i, 't_start'] = times[0]
                stim_df.loc[i, 't_stop'] = times[-1]
                signal = get_stimulus_signal(reader, exp, patterns[0], led=colors[0], returnas="percent")
                stim_df.loc[i, 'signal'] = signal
                stim_df.loc[i, 'intensity'] = max(signal).item()
            i += 1

        else:
            alignment = np.array(reader.electrode_intensity_alignment)
            if len(alignment.shape) != 0:
                exp_timer = round(alignment[0] * period, 5)

            for k, sub_stage in enumerate(substages):
                try:
                    sub_stage[0].name
                except:
                    sub_stage = sub_stage.copy_expand_ref()
                if (k != 0) and (len(alignment.shape) != 0):
                    times = [round(x * period, 5) for x in alignment]
                    stop_index = times.index(true_t_stop)
                    exp_timer = times[stop_index + 1]  # Grabbing one sample after the t_stop value for prior sub_stage.
                stim_df.loc[i + k, 'experiment'] = str(exp)
                stim_df.loc[i + k, 'stage'] = reader.experiment_stage_name
                stim_df.loc[i + k, 'substage'] = sub_stage.name

                patterns = [x.name for x in sub_stage.shapes]
                stim_df.loc[i + k, 'pattern'] = tuple(patterns)

                functions = [x.name for x in sub_stage.functions]
                stim_df.loc[i + k, 'function'] = functions

                color_mask = [sub_stage.color_r, sub_stage.color_g, sub_stage.color_b]
                colors = list(compress(led, color_mask))
                stim_df.loc[i + k, 'color'] = colors

                if len(alignment.shape) == 0:
                    stim_df.loc[i + k, 't_start'] = np.nan
                    stim_df.loc[i + k, 't_stop'] = np.nan
                    stim_df.loc[i + k, 'signal'] = [np.nan]
                    stim_df.loc[i + k, 'intensity'] = np.nan

                else:
                    duration = sub_stage.functions[0].duration * s
                    t_stop = exp_timer + duration
                    signal = get_stimulus_signal(reader, exp, patterns[0], led=colors[0], returnas="percent")
                    signal = signal.time_slice(t_start=exp_timer, t_stop=t_stop)
                    true_t_stop = round(signal.t_stop, 5)
                    stim_df.loc[i + k, 't_start'] = signal.t_start
                    stim_df.loc[i + k, 't_stop'] = true_t_stop
                    stim_df.loc[i + k, 'signal'] = signal
                    stim_df.loc[i + k, 'intensity'] = max(signal).item()
            i += k + 1
    stim_df.dropna(inplace=True, how='all')
    return stim_df


def write_exp_df_to_excel(exp_df, excel, sheet):
    book = load_workbook(excel)
    writer = pd.ExcelWriter(excel, engine='openpyxl')
    writer.book = book
    writer.sheets = dict((ws.title, ws) for ws in book.worksheets)
    exp_df.to_excel(writer, sheet_name=sheet, startrow=0, startcol=0, header=True, index=True)
    writer.save()


def read_exp_df_from_excel(excel, sheet):
    excel_file = pd.ExcelFile(excel)
    stim_df = excel_file.parse(sheet)
    return stim_df


def zero_runs(a):
    """
    Returns a Nx2 dimensional array, with each row containing the start and stop indices of contiguous zeros in the
    original array, a.
    """
    # Create an array that is 1 where a is 0, and pad each end with an extra 0.
    iszero = np.concatenate(([0], np.equal(a, 0).view(np.int8), [0]))
    absdiff = np.abs(np.diff(iszero))
    # Runs start and end where absdiff is 1.
    ranges = np.where(absdiff == 1)[0].reshape(-1, 2)
    return ranges


def divy_exp_series(exp_df, exps, sub_exps=10):

    new_exp_df = exp_df
    for exp in exps:
        if not isinstance(exp, str):
            exp = str(exp)
        sub_exp_df = pd.DataFrame(index=range(sub_exps), columns=exp_df.columns) #['experiment', 't_start', 't_stop', 'pattern', 'intensity', 'color', 'signal']
        sub_exp_df.loc[:, 'pattern'] = exp_df[exp_df['experiment']==exp]["pattern"].values[0] * sub_exps

        signal = exp_df[exp_df['experiment']==exp]["signal"].values
        signal = signal[0] #IrregularlySampledSignal object
        signal_aslist = signal.reshape(signal.shape[0]).tolist()
        zero_periods = zero_runs(signal_aslist)
        if zero_periods.shape[0] == sub_exps:
            zero_periods = np.insert(zero_periods, 0, [-1, 0], axis=0)

        if zero_periods.shape[0] != sub_exps+1:
            raise Exception("sub_exps it not equal to the number of intertrial intervals in the original data"
                            + str(zero_periods.shape[0]), str(zero_periods))

        for sub_exp in range(0, sub_exps):
            start_index = zero_periods[sub_exp, 1]
            stop_index = zero_periods[sub_exp+1, 0]
            sub_signal = signal[start_index:stop_index]

            sub_exp_df.loc[sub_exp, "t_start"] = sub_signal.t_start
            sub_exp_df.loc[sub_exp, "t_stop"] = sub_signal.t_stop
            sub_exp_df.loc[sub_exp, "intensity"] = round(max(sub_signal).item(), 2)
            sub_exp_df.loc[sub_exp, "experiment"] = exp + "-" + str(sub_exp)
            sub_exp_df.loc[sub_exp, "signal"] = sub_signal

        new_exp_df = new_exp_df.append(sub_exp_df)
        new_exp_df = new_exp_df[new_exp_df.experiment != exp]
    return new_exp_df


if __name__ == "__main__":

    #
    # date = "10-30-20"
    # slice = 3
    #
    # base_filename = r'D:\{date}\{file}.{ext}'
    # h5_file = date + "___slice" + str(slice) + "_merged"
    # h5_file = base_filename.format(date=date, file=h5_file, ext="h5")

    h5_file = 'D:\myRecordings\\1-14-21___slice2_merged.h5'

    reader = CeedDataReader(h5_file)
    reader.open_h5()
    reader.load_mcs_data()

    electrode = 'A4'
    offset, scale = reader.get_electrode_offset_scale(electrode)
    fs = reader.electrodes_metadata[electrode]['sampling_frequency'] * Hz
    period = 1. / fs
    raw_data = (np.array(reader.electrodes_data[electrode]) - offset) * scale
    raw_data = raw_data * 2000000
    raw_signal = neo.core.AnalogSignal(raw_data, sampling_rate=fs, units='uV')
    # Plot data from electrode of interest
    plt.plot(raw_signal.times, raw_signal, 'k', lw=0.5)

    # Grab stimulus signal for given experiment, and plot it alongside electrode data
    exp = 1
    stim = get_stimulus_signal(reader, exp, shape="enclosed")
    plt.plot(stim.times, stim, 'b', lw=2)
    plt.show()

    # Save experiment stimulus information to excel file
    exp_df = get_all_exps(h5_file)  # Grab all experiment information
    # exp_excel = base_filename.format(date=date, file="experiment_df", ext="xlsx")  # exp_excel is your excel filename
    exp_excel = 'D:\myRecordings\\experiment_df.xlsx'
    sheet = "slice" + str(slice)
    write_exp_df_to_excel(exp_df, exp_excel, sheet)  # write experiment information for given slice to exp_excel

    # Load experiment stimulus information from excel file
    exp_df = read_exp_df_from_excel(exp_excel, sheet)
    exp_df = exp_df[exp_df['experiment'] == exp]
    stim_start = exp_df['t_start'].values[0]
    stim_stop = exp_df['t_stop'].values[0]
    pattern = exp_df['pattern'].values[0]
    intensity = exp_df['intensity'].values[0]
    function = exp_df['function'].values[0]

    print("Experiment " + str(exp) + ", consisting of " + str(pattern) + ", was presented as a " + function +
          " with intensity = " + str(intensity) + ", from " + str(stim_start) + " to " + str(stim_stop))
